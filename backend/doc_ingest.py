"""Unstructured-document import (demo: drop a messy doc in, get a project draft out).

extract_text() handles pdf/docx/xlsx/csv/txt without ever calling a model.
ai_draft() structures the text via Requesty (strict JSON only).
regex_draft() is the offline fallback: finds endpoints/URLs/DB-like names,
never invents prose. All content stays fictional demo data.
"""
import csv
import io
import os
import re

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 12000

KINDS = {"api", "database", "repo", "dashboard", "widget", "computation",
         "report", "doc", "other"}
ENTRY_KINDS = {"definition", "caveat", "howto", "process", "note"}
DB_STOPLIST = {"API", "URL", "URI", "EOD", "CET", "IST", "PDF", "DOCX", "XLSX",
               "CSV", "KT", "SLA", "REQ", "JSON", "HTTP", "HTTPS", "WWW", "COM",
               "THE", "AND", "FOR", "WITH", "FROM", "THIS", "THAT", "NOTE",
               "EEPE", "PFE", "EAD", "CVAR", "CVA", "CCR", "RWA", "LR", "FILE"}

EXTRACT_PROMPT = (
    "You turn an unstructured team document into a project-registration draft. "
    "Reply with ONLY a JSON object, no fences, no commentary:\n"
    '{"project_name": str, "summary": str (2-3 sentences), "owner": str (person name or ""), '
    '"components": [{"kind": one of api|database|repo|dashboard|widget|computation|report|doc|other, '
    '"label": str, "ref": str, "notes": str}], '
    '"attributes": [{"key": str, "value": str}], '
    '"suggested_entries": [{"title": str, "body": str (2-5 sentences), '
    '"entry_kind": one of definition|caveat|howto|process|note}]}\n'
    "Rules: extract ONLY what the document states (endpoints, DB/table names, repo URLs, "
    "dashboards, widgets, people, cadences, cut-off times). Never invent refs. "
    "Max 12 components, 6 attributes, 5 entries. Unsure of a field: use \"\".\n"
    "DOCUMENT:\n\"\"\"\n{TEXT}\n\"\"\""
)


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, ValueError):
            continue
    return ""


def extract_text(filename: str, content: bytes) -> str:
    """Raw text out of any supported file. Raises ValueError when unsupported."""
    if len(content) > MAX_BYTES:
        raise ValueError("file too large (10MB max)")
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".txt", ".md", ".markdown", ".log", ""):
        text = _decode(content)
    elif ext == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(_decode(content))))
        except Exception:
            raise ValueError("could not parse CSV")
        text = "\n".join(" | ".join(c.strip() for c in r if c.strip()) for r in rows if any(c.strip() for c in r))
    elif ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ValueError("could not parse Excel file")
        parts = []
        for ws in wb.worksheets[:5]:
            parts.append(f"--- sheet: {ws.title} ---")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    break
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    elif ext == ".xls":
        raise ValueError("legacy .xls not supported — re-save as .xlsx or .csv")
    elif ext == ".docx":
        from docx import Document as DocxDocument
        try:
            doc = DocxDocument(io.BytesIO(content))
        except Exception:
            raise ValueError("could not parse Word file")
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    elif ext == ".doc":
        raise ValueError("legacy .doc not supported — re-save as .docx")
    elif ext == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception:
            raise ValueError("could not parse PDF")
        parts = []
        for page in reader.pages[:20]:
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                parts.append(t)
        text = "\n".join(parts)
        if not text.strip():
            raise ValueError("PDF has no readable text (scanned image?)")
    else:
        text = _decode(content)
        if not text.strip() or len(re.findall(r"[^\x09\x0a\x0d\x20-\x7e]", text)) > len(text) // 4:
            raise ValueError(f"unsupported file type ({ext or 'unknown'}) — use pdf/docx/xlsx/csv/txt")
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not text:
        raise ValueError("no readable text found in file")
    return text[:MAX_CHARS]


def extract_json_block(text: str):
    """First balanced {...} span (string-aware). None if absent."""
    import json as _json
    start = text.find("{")
    if start < 0:
        return None
    depth, instr, esc = 0, False, False
    for i in range(start, len(text)):
        ch = text[i]
        if instr:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                instr = False
        elif ch == '"':
            instr = True
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return text[start:i + 1]
    return None


def parse_draft_json(text: str):
    """Tolerant parse: balanced span, trailing commas stripped, lax strings."""
    import json as _json
    block = extract_json_block(text)
    if not block:
        return None
    block = re.sub(r",\s*([}\]])", r"\1", block)
    try:
        return _json.loads(block, strict=False)
    except ValueError:
        return None


_DRAFT_KEYS = {"project_name", "name", "summary", "owner", "components",
               "attributes", "suggested_entries", "draft_project"}


def pick_draft_json(text: str):
    """All balanced top-level {...} blocks; return the one looking most like
    a project draft (thinking traces often contain small inline fragments)."""
    import json as _json
    cands = []
    i, n = 0, len(text)
    while True:
        j = text.find("{", i)
        if j < 0:
            break
        depth, instr, esc = 0, False, False
        k = j
        while k < n:
            ch = text[k]
            if instr:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    instr = False
            elif ch == '"':
                instr = True
            elif ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    break
            k += 1
        if depth == 0 and k < n:
            frag = text[j:k + 1]
            try:
                obj = _json.loads(re.sub(r",\s*([}\]])", r"\1", frag), strict=False)
                if isinstance(obj, dict):
                    score = len(set(obj.keys()) & _DRAFT_KEYS)
                    cands.append((score, len(frag), obj))
            except ValueError:
                pass
            i = k + 1
        else:
            i = j + 1
    if not cands:
        return None
    cands.sort(key=lambda t: (t[0], t[1]), reverse=True)
    return cands[0][2] if cands[0][0] > 0 else None


def sanitize_draft(raw: dict, filename: str) -> dict:
    """Coerce model output into the draft shape; unknown kinds fall back safely."""
    if not isinstance(raw, dict):
        raise ValueError("model did not return a JSON object")
    comps = []
    for c in (raw.get("components") or [])[:12]:
        if not isinstance(c, dict):
            continue
        comps.append({"kind": c.get("kind") if c.get("kind") in KINDS else "other",
                      "label": str(c.get("label", ""))[:120],
                      "ref": str(c.get("ref", ""))[:300],
                      "notes": str(c.get("notes", ""))[:300]})
    attrs = []
    for a in (raw.get("attributes") or [])[:6]:
        if isinstance(a, dict) and str(a.get("key", "")).strip():
            attrs.append({"key": str(a["key"])[:80], "value": str(a.get("value", ""))[:300]})
    entries = []
    for e in (raw.get("suggested_entries") or [])[:5]:
        if not isinstance(e, dict) or not str(e.get("title", "")).strip():
            continue
        entries.append({"title": str(e["title"])[:120], "body": str(e.get("body", ""))[:2000],
                        "entry_kind": e.get("entry_kind") if e.get("entry_kind") in ENTRY_KINDS else "note"})
    proj = raw.get("draft_project", raw) if isinstance(raw.get("draft_project"), dict) else raw
    return {"draft_project": {"name": str(proj.get("project_name", proj.get("name", "")))[:80],
                             "summary": str(proj.get("summary", ""))[:500],
                             "owner": str(proj.get("owner", ""))[:80]},
            "components": comps, "attributes": attrs, "suggested_entries": entries}


def regex_draft(text: str, filename: str) -> dict:
    """Offline fallback: find machine-shaped strings, invent nothing else."""
    raw_urls = sorted(set(re.findall(r"https?://\S+|(?:github\.com|appstudio://|sg-dash://)\S*", text)))[:12]
    urls = [u.rstrip(".,;)") for u in raw_urls]
    blob = " ".join(urls)
    endpoints = sorted({e.rstrip(".,;)") for e in re.findall(r"/[A-Za-z0-9][A-Za-z0-9/_\-.{}]*", text)
                        if e.rstrip(".,;)") not in blob})[:12]
    dbs = sorted({t for t in re.findall(r"\b[A-Z][A-Z0-9_]{3,}\b", text) if t not in DB_STOPLIST})[:10]
    comps = [{"kind": "api", "label": e, "ref": e, "notes": "found in " + filename} for e in endpoints]
    for u in urls:
        kind = "repo" if "github" in u else ("report" if u.startswith("appstudio") else
                                            ("dashboard" if u.startswith("sg-dash") else "doc"))
        comps.append({"kind": kind, "label": u.split("/")[-1][:60] or u[:60], "ref": u,
                      "notes": "found in " + filename})
    comps += [{"kind": "database", "label": d, "ref": d, "notes": "found in " + filename} for d in dbs]
    stem = os.path.splitext(os.path.basename(filename or "upload"))[0].replace("_", " ").replace("-", " ")
    first_lines = " ".join(t for t in text.split("\n") if t.strip())[:300]
    return {"draft_project": {"name": stem[:80], "summary": first_lines, "owner": ""},
            "components": comps[:12], "attributes": [], "suggested_entries": []}
